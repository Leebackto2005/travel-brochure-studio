"""Native editable travel brochure toolkit. No network or generative calls."""
import argparse, copy, hashlib, json, math, re, sys, zipfile
from pathlib import Path
from lxml import etree as ET
from PIL import Image, ImageDraw, ImageFont, ImageChops
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
STYLES = {
 'editorial':dict(name='旅行编辑',ink='1C303A',primary='1C303A',accent='EA603D',paper='FAF9F5',pale='E7EEEE'),
 'zine':dict(name='旅行手记',ink='234B40',primary='234B40',accent='D56B43',paper='FAF3E8',pale='E3E9DC'),
 'swiss':dict(name='瑞士信息网格',ink='181E20',primary='2140EA',accent='F05E42',paper='FFFFFF',pale='EBEFFF'),
 'boutique':dict(name='精品旅行',ink='29453D',primary='29453D',accent='A98B5C',paper='F3F0E7',pale='E0E5DA')}
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
def dump(p,d):
 p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
def read(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def fail(s): raise ValueError(s)
def project(p):
 d=read(p); d['_base']=str(Path(p).resolve().parent); return d
def asset_path(d,a): return (Path(d['_base'])/a['path']).resolve() if a.get('path') else None
def walk_shapes(shapes):
 for s in shapes:
  yield s
  if s.shape_type==6: yield from walk_shapes(s.shapes)
def numeric(s):
 s=s.upper().replace('公斤','KG').replace('千克','KG').replace('元','CNY')
 return sorted(re.findall(r'\b[A-Z]{2}\d{3,4}\b|\d{1,2}:\d{2}|\d+(?:\.\d+)?\s*KG|CNY\s*\d+|\d+\s*CNY',s))
def validate(d):
 errors=[]; warnings=[]
 if d.get('schema')!='travel-brochure/v1': errors.append('schema must be travel-brochure/v1')
 if d.get('style') not in STYLES: errors.append('Unknown style')
 if not d.get('title'): errors.append('title is required')
 assets=d.get('assets',[]); ids=[a.get('id') for a in assets]
 if len(ids)!=len(set(ids)) or None in ids: errors.append('Asset IDs must be unique and nonempty')
 for a in assets:
  p=asset_path(d,a)
  if p:
   if not p.is_file(): errors.append(f'Missing asset: {p}')
   else:
    try:
     with Image.open(p) as im: im.verify()
    except Exception: errors.append(f'Unreadable image: {p}')
  else: warnings.append(f'Placeholder: {a.get("id")}')
  if not a.get('source'): errors.append(f'Asset source missing: {a.get("id")}')
  if a.get('fit','cover') not in ('cover','contain'): errors.append('fit must be cover or contain')
  f=a.get('focus',[.5,.5])
  if not isinstance(f,list) or len(f)!=2 or any(not isinstance(v,(int,float)) or not 0<=v<=1 for v in f): errors.append('focus must contain two values from 0 to 1')
 def refs(x):
  if isinstance(x,dict):
   for k,v in x.items():
    if k in ('asset','secondary_asset','exterior','room') and v and v not in ids: errors.append(f'Unknown asset reference: {v}')
    refs(v)
   if 'zh' in x and 'en' in x and numeric(x['zh'])!=numeric(x['en']): errors.append(f'Bilingual numeric mismatch: {x["zh"][:50]}')
  elif isinstance(x,list):
   for v in x: refs(v)
 pages=d.get('pages',[])
 if not pages: errors.append('pages is empty')
 for i,p in enumerate(pages,1):
  typ=p.get('type'); refs(p)
  if not p.get('title'): errors.append(f'Page {i}: title required')
  cap={'cover':0,'overview':5,'gallery':6,'meals':4,'services':4,'hotels':3,'itinerary':2,'terms':6}
  if typ not in cap: errors.append(f'Page {i}: unknown page type'); continue
  key='days' if typ=='itinerary' else 'rows' if typ=='terms' else 'items'
  if typ!='cover' and not 1<=len(p.get(key,[]))<=cap[typ]: errors.append(f'Page {i}: {key} capacity 1..{cap[typ]}')
  for item in p.get(key,[]):
   if typ=='itinerary':
    for k in ('day','title','en_title','zh','en'):
     if not item.get(k): errors.append(f'Page {i}: day.{k} required')
    if len(item.get('zh',''))>210 or len(item.get('en',''))>680: errors.append(f'Page {i}: split long itinerary, do not shrink text')
   elif typ=='terms':
    if not all(k in item for k in ('label','zh','en')): errors.append(f'Page {i}: terms require label/zh/en')
    if len(item.get('en',''))>250 or len(item.get('zh',''))>110: errors.append(f'Page {i}: split long terms')
   elif not item.get('title'): errors.append(f'Page {i}: item.title required')
  if typ=='hotels':
   for item in p.get('items',[]):
    if not item.get('exterior') or not item.get('room'): errors.append(f'Page {i}: hotel must pair exterior and room')
 return {'errors':errors,'warnings':warnings,'semantic_review':'required'}

def specimen(style):
 assets=[dict(id=x,path=None,label=label,source='待补充真实素材',verified=False,fit='cover',focus=[.5,.5]) for x,label in [('hero','目的地主视觉'),('detail','旅行细节'),('a','主推景点 A'),('b','主推景点 B'),('c','文化体验'),('d','自然风景'),('ext','酒店外观'),('room','双床房内饰')]]
 items=[dict(title=z['label'],en='Destination photograph',asset=z['id'],note='照片与产品内容核对后替换') for z in assets[2:6]]
 return dict(schema='travel-brochure/v1',title='海岛慢旅行',style=style,sources=[],assets=assets,pages=[
  dict(type='cover',title='海岛慢旅行',en='A JOURNEY, DAY BY DAY',asset='hero',secondary_asset='detail',highlights=['可编辑示例 · 非正式产品','按实际行程填写天数和出发信息','原生文字 / 真实照片 / 双语对照']),
  dict(type='overview',title='把旅程写进日常',en='JOURNEY AT A GLANCE',items=items),
  dict(type='gallery',title='精选旅行体验',en='SELECTED EXPERIENCES',paired_featured=True,items=items),
  dict(type='hotels',title='休息，也是风景',en='REST & STAY',items=[dict(title='备选酒店名称',en='Hotel option — to be confirmed',exterior='ext',room='room',room_label='外观与双床房须为同一家酒店 / Matching exterior and twin room'),dict(title='第二家备选酒店',en='Second hotel option',exterior='ext',room='room',room_label='实际房型及入住安排以最终确认为准 / Subject to confirmation')]),
  dict(type='itinerary',title='每日行程',en='THE DAILY PLAN',days=[dict(day=i,title=t,en_title=e,zh='这里放置已确认的中文行程。保留景点顺序、用餐安排与必要条件；英文紧随其下。',en='Place the confirmed itinerary here. Preserve the order of visits, meal arrangements and all applicable conditions.',asset='a' if i==1 else 'b',flight='航班信息待填 / Flight details to be confirmed',meal_zh='餐食按最终行程安排',meal_en='Meals as specified in the confirmed itinerary.',hotel_zh='住宿待确认',hotel_en='Accommodation to be confirmed.') for i,t,e in [(1,'抵达目的地','Arrival'),(2,'探索与体验','Explore & discover')]]),
  dict(type='terms',title='出行须知',en='TRAVEL ESSENTIALS',rows=[dict(label='费用包含 / Inclusions',zh='以最终确认的产品内容为准。',en='Subject to the confirmed package details.'),dict(label='自费项目 / Optional activities',zh='自愿参加，费用自理。',en='Participation is voluntary and at your own expense.'),dict(label='住宿 / Accommodation',zh='酒店及房型以最终确认为准。',en='Hotels and room types are subject to final confirmation.')],notice_zh='行程内容若有分歧，以中文版为准。',notice_en='In case of any discrepancy between the Chinese and English versions of the itinerary, the Chinese version shall prevail.')])

class Painter:
 def __init__(self,d):
  self.d=d;self.style=d['style'];self.c=STYLES[self.style]|d.get('style_overrides',{});self.prs=Presentation();self.prs.slide_width=Inches(8.2677);self.prs.slide_height=Inches(11.6929);self.assets={a['id']:a for a in d['assets']};self.log=[]
 def rect(self,x,y,w,h,color,round=False):
  s=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color);s.line.fill.background();s.shadow.inherit=False;return s
 def text(self,t,x,y,w,h,size=12,color=None,bold=False,serif=False):
  s=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  for n,line in enumerate(str(t).split('\n')):
   p=tf.paragraphs[0] if n==0 else tf.add_paragraph();p.text=line;p.font.name='Georgia' if serif else 'Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color or self.c['ink']);p.space_after=Pt(4)
   for r in p.runs:
    ea=OxmlElement('a:ea');ea.set('typeface','SimSun' if serif else 'Microsoft YaHei');r._r.get_or_add_rPr().append(ea)
  return s
 def line(self,x,y,w,color=None): self.rect(x,y,w,.009,color or self.c['ink'])
 def photo(self,aid,x,y,w,h):
  a=self.assets.get(aid);path=asset_path(self.d,a) if a else None
  self.rect(x,y,w,h,self.c['pale'])
  if not path:
   self.text((a or {}).get('label','照片待补充'),x+.12,y+h/2-.15,w-.24,.5,11,self.c['primary']);return
  with Image.open(path) as im: iw,ih=im.size
  fit=a.get('fit','cover');fx,fy=a.get('focus',[.5,.5]);src=iw/ih;box=w/h
  if fit=='contain':
   nw,nh=(w,w/src) if src>box else (h*src,h)
   pic=self.s.shapes.add_picture(str(path),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))
  else:
   pic=self.s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
   if src>box: crop=1-box/src;pic.crop_left=crop*fx;pic.crop_right=crop*(1-fx)
   else: crop=1-src/box;pic.crop_top=crop*fy;pic.crop_bottom=crop*(1-fy)
  pic.name=f'asset:{aid}';self.log.append(dict(asset=aid,page=len(self.prs.slides),width_px=iw,height_px=ih,fit=fit,focus=[fx,fy],effective_ppi=round(min(iw/w,ih/h)),source=a.get('source'),verified=a.get('verified',False)))
 def header(self,p,n):
  c=self.c;s=self.style
  if s=='editorial': self.rect(0,0,.085,11.7,c['accent']);self.text('TRAVEL / EDITION',.45,.35,5,.25,10,bold=True);self.line(.45,.73,7.36)
  elif s=='zine': self.rect(0,0,8.27,.12,c['primary']);self.text('JOURNEY / 旅行手记',.45,.35,6,.4,14,c['primary'],serif=True)
  elif s=='swiss': self.rect(0,0,8.27,.15,c['primary']);self.text('TRIP / ITINERARY',.4,.35,6,.35,14,bold=True);self.rect(.4,.85,2.0,.06,c['accent'])
  else: self.rect(0,0,8.27,.22,c['primary']);self.text('T H E  J O U R N E Y',.5,.45,6,.3,11,serif=True);self.line(.5,.88,7.27,c['accent'])
  self.text(p['title'],.45,1.03,7.35,.65,27,bold=s=='swiss',serif=s in ('editorial','boutique'))
  self.text(p.get('en',''),.48,1.77,7.2,.35,11,c['accent'],True)
  self.line(.45,11.12,7.35);self.text('TRAVEL COLLECTION',.45,11.25,5,.2,8);self.text(f'{n:02d} / {len(self.d["pages"]):02d}',7.08,11.25,.75,.2,8)
 def caption(self,item,x,y,w):
  self.text(item['title'],x,y,w,.28,12,bold=True);self.text(item.get('en',''),x,y+.32,w,.43,9)
 def cover(self,p):
  s=self.style;c=self.c
  if s=='editorial':
   self.photo(p.get('asset'),.45,2.5,5.0,5.1);self.photo(p.get('secondary_asset'),5.65,4.65,2.17,2.95);self.text('旅程概要\nOVERVIEW',5.65,2.5,2.17,1.0,17,bold=True)
  elif s=='zine':
   self.photo(p.get('asset'),.45,2.5,7.37,3.5);self.photo(p.get('secondary_asset'),.65,6.25,2.5,2.8);self.rect(3.38,6.25,4.44,.58,c['primary']);self.text('ROUTE / 让旅行慢下来',3.55,6.42,4,.3,12,'FFFFFF',True)
  elif s=='swiss':
   self.rect(.45,2.45,1.15,5.5,c['primary']);self.text('01\n→',.6,2.8,.9,1.8,35,'FFFFFF',True);self.photo(p.get('asset'),1.8,2.45,6.02,3.4);self.photo(p.get('secondary_asset'),5.35,6.05,2.47,1.9);self.text('DESTINATION\n旅行，从这里开始',1.8,6.25,3.2,1.2,18,bold=True)
  else:
   self.photo(p.get('asset'),.45,2.65,4.6,4.8);self.photo(p.get('secondary_asset'),5.25,2.65,2.57,4.8);self.line(.45,7.8,7.37,c['accent']);self.text('A CURATED JOURNEY',.45,8.03,7,.5,20,c['accent'],serif=True)
  y=9.35 if s=='zine' else 8.8
  for i,t in enumerate(p.get('highlights',[])[:3]): self.text(t,.5,y+i*.48,7.25,.4,12)
 def overview(self,p):
  items=p['items'];s=self.style
  if s=='boutique':
   self.line(.75,2.6,6.7,self.c['accent'])
   for i in range(len(items)): self.text(f'{i+1:02d}',.65+i*6.7/max(1,len(items)-1),2.25,.6,.3,12,self.c['accent'])
  top=3.1 if s=='boutique' else 2.5;rh=(10.9-top)/len(items)
  for i,it in enumerate(items):
   y=top+i*rh;self.line(.45,y+rh-.12,7.37,self.c['pale'])
   if s=='zine' and i%2:
    self.text(f'{i+1:02d}',7.12,y,.7,.5,26,self.c['accent'],serif=True);self.photo(it.get('asset'),4.65,y,2.25,rh-.32);self.caption(it,.6,y+.08,3.75)
   elif s=='swiss':
    self.rect(.45,y,.85,rh-.25,self.c['primary'] if i%2==0 else self.c['ink']);self.text(f'{i+1:02d}',.55,y+.15,.7,.5,25,'FFFFFF',True);self.photo(it.get('asset'),1.48,y,2.1,rh-.25);self.caption(it,3.8,y+.08,3.8)
   else:
    self.text(f'{i+1:02d}',.45,y,.7,.6,28,self.c['accent'],serif=True);self.photo(it.get('asset'),1.28,y,2.4,rh-.28);self.caption(it,3.95,y+.1,3.85)
 def gallery(self,p):
  items=p['items'];n=len(items);top=2.5
  if p.get('paired_featured') and n>=2:
   for i,it in enumerate(items[:2]):
    x=.45+i*3.8;self.photo(it.get('asset'),x,top,3.57,3.5);self.caption(it,x,6.16,3.57)
   tail=items[2:];w=(7.37-.18*max(0,len(tail)-1))/max(1,len(tail))
   for i,it in enumerate(tail): x=.45+i*(w+.18);self.photo(it.get('asset'),x,7.4,w,2.3);self.caption(it,x,9.86,w)
  elif p.get('featured') and n>1:
   self.photo(items[0].get('asset'),.45,top,7.37,3.5);self.caption(items[0],.45,6.12,7.37);w=(7.37-.2*(n-2))/(n-1)
   for i,it in enumerate(items[1:]): x=.45+i*(w+.2);self.photo(it.get('asset'),x,7.3,w,2.2);self.caption(it,x,9.64,w)
  else:
   cols=2;rows=math.ceil(n/2);rh=8.25/rows
   for i,it in enumerate(items):
    x=.45+(i%2)*3.8;y=top+(i//2)*rh;self.photo(it.get('asset'),x,y,3.57,rh-.95);self.caption(it,x,y+rh-.8,3.57)
 def hotels(self,p):
  rh=8.35/len(p['items'])
  for i,it in enumerate(p['items']):
   y=2.45+i*rh;self.text(it['title'],.45,y,7.3,.32,15,bold=True);self.text(it.get('en',''),.45,y+.36,7.3,.3,9)
   ph=min(2.7,rh-1.25);self.photo(it.get('exterior'),.45,y+.78,3.57,ph);self.photo(it.get('room'),4.25,y+.78,3.57,ph);self.text(it.get('room_label',''),.45,y+.88+ph,7.37,.44,9)
 def itinerary(self,p):
  days=p['days'];rh=8.4/len(days)
  for i,it in enumerate(days):
   y=2.45+i*rh;sw=self.style=='swiss';self.rect(.45,y,.75,.66,self.c['primary']);self.text(f'{it["day"]:02d}',.56,y+.12,.65,.45,24,'FFFFFF',True)
   self.text(it['title'],1.4,y,6.3,.35,17,bold=True);self.text(it['en_title'],1.4,y+.4,6.3,.3,10,self.c['accent'])
   self.text(it.get('flight',''),.48,y+.88,7.3,.35,10,self.c['primary'],True)
   self.photo(it.get('asset'),.45,y+1.38,1.65,1.58)
   self.text(it['zh'],2.33,y+1.33,5.49,1.05,11)
   self.text(it['en'],2.33,y+2.44,5.49,1.05,10)
   extras=' / '.join(filter(None,[it.get('meal_zh'),it.get('hotel_zh')]))+'\n'+' / '.join(filter(None,[it.get('meal_en'),it.get('hotel_en')]))
   self.text(extras,.48,y+3.55,7.3,.48,8.5);self.line(.45,y+rh-.06,7.37,self.c['pale'])
 def terms(self,p):
  y=2.55;rh=min(1.05,6.65/max(1,len(p['rows'])))
  for it in p['rows']:
   self.rect(.45,y,1.75,rh-.03,self.c['primary']);self.text(it['label'],.58,y+.13,1.48,rh-.19,11,'FFFFFF',True);self.rect(2.2,y,5.62,rh-.03,self.c['pale']);self.text(it['zh'],2.35,y+.1,5.3,.38,10.5);self.text(it['en'],2.35,y+.5,5.3,rh-.55,9.5);y+=rh
  self.line(.45,9.15,7.37,self.c['accent']);self.text(p.get('notice_zh',''),.5,9.45,7.2,.4,12,bold=True);self.text(p.get('notice_en',''),.5,10.0,7.2,.7,10)
 def build(self,out):
  for n,p in enumerate(self.d['pages'],1):
   self.s=self.prs.slides.add_slide(self.prs.slide_layouts[6]);self.s.background.fill.solid();self.s.background.fill.fore_color.rgb=RGBColor.from_string(self.c['paper']);self.header(p,n)
   getattr(self, 'gallery' if p['type'] in ('gallery','meals','services') else p['type'])(p)
  self.prs.save(out);dump(str(out)+'.assets.json',self.log)

def plan(path,out):
 d=project(path);r=validate(d)
 if r['errors']: fail('\n'.join(r['errors']))
 manifest={'project_sha256':sha(path),'style':d['style'],'pages':[{'page':i,'type':p['type'],'title':p['title']} for i,p in enumerate(d['pages'],1)],'assets':[]}
 for a in d['assets']:
  q=asset_path(d,a);entry={'id':a['id'],'source':a.get('source'),'path':str(q) if q else None}
  if q:
   with Image.open(q) as im: entry.update(size=list(im.size),sha256=sha(q))
  manifest['assets'].append(entry)
 dump(Path(out)/'plan.json',manifest);dump(Path(path).with_suffix('.plan.json'),manifest)
 Path(out).mkdir(parents=True,exist_ok=True)
 (Path(out)/'Design-Bible.md').write_text('# 视觉设计底稿\n\n风格：'+STYLES[d['style']]['name']+'\n\nA4竖版；原生文字；按主体焦点裁图；双语上中下英。\n\n需在代表页审阅时补充：受众、视觉理由、主推优先级、不可裁切主体、实际审阅结论。\n\n'+ '\n'.join(f'- P{x["page"]}: {x["title"]} / {x["type"]}' for x in manifest['pages']),encoding='utf-8')
 return manifest

def inspect(path,out):
 path=Path(path);out=Path(out);out.mkdir(parents=True,exist_ok=True);result={'source':str(path.resolve()),'sha256':sha(path),'pages':[],'warning':'Extracted text may be hidden by pictures; check rendered pages.'}
 if path.suffix.lower()=='.pptx':
  p=Presentation(path)
  for i,slide in enumerate(p.slides,1):
   shapes=[]
   for s in walk_shapes(slide.shapes):
    q={'id':s.shape_id,'name':s.name,'bbox_inches':[round(v/914400,3) for v in (s.left,s.top,s.width,s.height)]}
    if s.has_text_frame:q['text']=s.text
    if s.has_table:q['table']=[[c.text for c in r.cells] for r in s.table.rows]
    if s.shape_type==13:
     f=f'P{i:02d}-{s.shape_id}.{s.image.ext}';(out/f).write_bytes(s.image.blob);q['image']=f
    shapes.append(q)
   result['pages'].append({'page':i,'shapes':shapes})
 elif path.suffix.lower()=='.pdf':
  import fitz
  with fitz.open(path) as pdf:
   for i,p in enumerate(pdf,1):
    result['pages'].append({'page':i,'text':p.get_text()});p.get_pixmap(matrix=fitz.Matrix(1,1)).save(out/f'P{i:02d}.png')
 elif path.suffix.lower()=='.docx':
  with zipfile.ZipFile(path) as z:
   root=ET.fromstring(z.read('word/document.xml'));n={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};result['paragraphs']=[''.join(p.itertext()) for p in []];result['paragraphs']=[''.join(p.xpath('.//w:t/text()',namespaces=n)) for p in root.xpath('//w:p',namespaces=n)]
 else:fail('Supported inputs: PPTX, PDF, DOCX')
 dump(out/'inventory.json',result);contact(out,out/'contact-sheet.jpg');return result
def contact(folder,out):
 paths=sorted(p for p in Path(folder).glob('*') if p.suffix.lower() in ('.png','.jpg','.jpeg') and p.name!=Path(out).name)
 if not paths:return
 font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',14) if Path('C:/Windows/Fonts/msyh.ttc').exists() else ImageFont.load_default()
 sheet=Image.new('RGB',(960,math.ceil(len(paths)/4)*260),'#eeeeee');draw=ImageDraw.Draw(sheet)
 for i,p in enumerate(paths):
  with Image.open(p) as im: im=im.convert('RGB');im.thumbnail((226,225));x=(i%4)*240+(240-im.width)//2;y=(i//4)*260;sheet.paste(im,(x,y));draw.text(((i%4)*240+8,y+230),p.stem[:22],fill='black',font=font)
 sheet.save(out)

def patch(source,changes,out):
 if Path(source).resolve()==Path(out).resolve():fail('Do not overwrite source')
 edits=read(changes)['changes'];changed={}
 with zipfile.ZipFile(source) as z:
  for edit in edits:
   member=f'ppt/slides/slide{int(edit["slide"])}.xml';root=ET.fromstring(changed.get(member,z.read(member)));sid=str(edit['shape_id']);old=edit['old'];new=edit['new']
   if not old or '\n' in old+new or '\r' in old+new:fail('Nonempty single-paragraph replacement required')
   shape=root.xpath(f'//p:sp[p:nvSpPr/p:cNvPr[@id="{sid}"]]',namespaces=NS)
   if len(shape)!=1:fail('Target shape not found or not unique')
   matches=[]
   for par in shape[0].xpath('.//a:p',namespaces=NS):
    ts=par.xpath('.//a:t',namespaces=NS);full=''.join(t.text or '' for t in ts)
    for m in re.finditer(re.escape(old),full):matches.append((ts,m.start(),m.end()))
   if len(matches)!=1:fail('Exact old text must occur exactly once in target shape')
   ts,start,end=matches[0];pos=0;inserted=False
   for t in ts:
    val=t.text or '';a=pos;b=pos+len(val);pos=b
    if b<=start or a>=end:continue
    prefix=val[:max(0,start-a)];suffix=val[max(0,end-a):] if end<b else ''
    t.text=prefix+(new if not inserted else '')+suffix;inserted=True
   changed[member]=ET.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
  Path(out).parent.mkdir(parents=True,exist_ok=True)
  with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as target:
   for info in z.infolist():target.writestr(info,changed.get(info.filename,z.read(info.filename)))
 dump(str(out)+'.patch.json',{'source_sha256':sha(source),'changed_members':list(changed),'changes':edits,'render_review':'required'})
def compare(before,after,allowed):
 aa={p.name:p for p in Path(before).glob('P[0-9]*.png')};bb={p.name:p for p in Path(after).glob('P[0-9]*.png')};changed=[];errors=[]
 if not aa or aa.keys()!=bb.keys():errors.append('Missing or unequal rendered page sets')
 for name in aa.keys()&bb.keys():
  with Image.open(aa[name]) as a,Image.open(bb[name]) as b:
   if a.size!=b.size or ImageChops.difference(a.convert('RGB'),b.convert('RGB')).getbbox():changed.append(name)
 for n in changed:
  if int(re.search(r'\d+',n)[0]) not in allowed:errors.append('Unauthorized page change: '+n)
 return dict(changed=changed,errors=errors)
def verify(ppt,render):
 p=Presentation(ppt);errors=[];r=Path(render)
 for i,slide in enumerate(p.slides,1):
  for s in slide.shapes:
   if min(s.left,s.top)<-2000 or s.left+s.width>p.slide_width+16000 or s.top+s.height>p.slide_height+16000:errors.append(f'P{i}: outside canvas: {s.name}')
 rendered=len(list(r.glob('P[0-9]*.png')))==len(p.slides)
 report=read(r/'render-report.json') if (r/'render-report.json').exists() else {}
 if not rendered:errors.append('Missing rendered PNG pages')
 if report.get('source_sha256','').lower()!=sha(ppt):errors.append('Missing/stale render report source hash')
 errors.extend(report.get('overflow',[]))
 pdf=r/'brochure.pdf';pdf_pages=None
 if pdf.exists():
  import fitz
  with fitz.open(pdf) as doc:pdf_pages=len(doc)
 if pdf_pages!=len(p.slides):errors.append('PDF count missing/mismatched')
 return dict(slides=len(p.slides),errors=errors,rendered=rendered,pdf_pages=pdf_pages,visual_review='required',business_review='required')

def main():
 ap=argparse.ArgumentParser(description=__doc__);sub=ap.add_subparsers(dest='cmd',required=True)
 q=sub.add_parser('init');q.add_argument('--style',choices=STYLES,default='editorial');q.add_argument('--out',required=True)
 for cmd in ('validate','plan','build'):
  q=sub.add_parser(cmd);q.add_argument('project')
  if cmd!='validate':q.add_argument('--out',required=True)
  if cmd=='build':q.add_argument('--allow-reflow',action='store_true',help='Use only when user explicitly authorizes a new layout')
 q=sub.add_parser('inspect');q.add_argument('source');q.add_argument('--out',required=True)
 q=sub.add_parser('patch');q.add_argument('source');q.add_argument('--changes',required=True);q.add_argument('--out',required=True)
 q=sub.add_parser('verify');q.add_argument('source');q.add_argument('--render-dir',required=True);q.add_argument('--out',required=True)
 q=sub.add_parser('compare');q.add_argument('before');q.add_argument('after');q.add_argument('--allowed',default='');q.add_argument('--out',required=True)
 q=sub.add_parser('contact');q.add_argument('folder');q.add_argument('--out',required=True)
 args=ap.parse_args();cmd=args.cmd
 if cmd=='init':dump(args.out,specimen(args.style))
 elif cmd=='validate':
  r=validate(project(args.project));print(json.dumps(r,ensure_ascii=False,indent=2));return bool(r['errors'])
 elif cmd=='plan':plan(args.project,args.out)
 elif cmd=='build':
  if not args.allow_reflow:fail('Strict template filling is default. Use strict_template.py; --allow-reflow requires explicit user intent.')
  d=project(args.project);r=validate(d)
  if r['errors']:fail('\n'.join(r['errors']))
  pp=Path(args.project).with_suffix('.plan.json')
  if not pp.exists() or read(pp)['project_sha256']!=sha(args.project):fail('Run plan for this exact project before build')
  for a in read(pp)['assets']:
   if a['path'] and sha(a['path'])!=a['sha256']:fail('Asset changed; inspect and re-plan')
  Path(args.out).parent.mkdir(parents=True,exist_ok=True);Painter(d).build(args.out)
 elif cmd=='inspect':inspect(args.source,args.out)
 elif cmd=='patch':patch(args.source,args.changes,args.out)
 elif cmd=='contact':contact(args.folder,args.out)
 else:
  r=verify(args.source,args.render_dir) if cmd=='verify' else compare(args.before,args.after,{int(x) for x in args.allowed.split(',') if x});dump(args.out,r);print(json.dumps(r,ensure_ascii=False,indent=2));return bool(r['errors'])
 print(f'{cmd}: complete');return 0
if __name__=='__main__':
 try:sys.exit(main())
 except (ValueError,KeyError,FileNotFoundError) as e:print(f'ERROR: {e}',file=sys.stderr);sys.exit(2)
