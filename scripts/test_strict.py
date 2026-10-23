import copy,tempfile,sys
from pathlib import Path
from lxml import etree as E
from PIL import Image
from pptx import Presentation
import strict_template as t,studio as s
def main():
 passed=[];root=s.ROOT
 with tempfile.TemporaryDirectory(prefix='strict-tests-') as td:
  td=Path(td);Image.new('RGB',(1200,400),'#6699AA').save(td/'photo.png')
  for svg in (root/'assets/reference-shapes').glob('*.svg'):
   c=t.catalog(svg);blank=td/'blank.pptx';t.build(svg,{},blank);p=Presentation(blank);scale=841.89/c['height']
   nodes=list(E.parse(str(svg)).getroot());shapes={z.name:z for z in p.slides[0].shapes}
   for i,e in enumerate(nodes):
    if t.tag(e)=='rect':
     z=shapes[f'svg-element-{i}'];assert all(abs(a-b)<3 for a,b in zip((z.left,z.top,z.width,z.height),(int(t.val(e,k)*scale*t.EMU) for k in ('x','y','width','height'))))
   assert len([z for z in p.slides[0].shapes if z.has_text_frame and z.name.startswith('text-')])==len(c['texts'])
   passed.append(svg.stem+': original rectangle geometry and editable texts preserved')
   data={'_base':str(td),'template_sha256':c['sha256'],'photos':{'photo-01':{'path':'photo.png'}},'bodies':{'body-01':'示例正文'},'texts':{}}
   output=td/'filled.pptx';t.build(svg,data,output);p2=Presentation(output);pic=next(z for z in p2.slides[0].shapes if z.shape_type==13);b=c['photos'][0]
   assert all(abs(a-b)<3 for a,b in zip((pic.left,pic.top,pic.width,pic.height),(int(b[k]*scale*t.EMU) for k in ('x','y','width','height'))))
   assert 0<=pic.crop_left<1;passed.append(svg.stem+': image stays inside original frame')
   long=copy.deepcopy(data);long['bodies']['body-01']='不能自动缩字。'*200
   try:t.build(svg,long,output);raise AssertionError('Long copy accepted')
   except ValueError:passed.append(svg.stem+': overflow rejected')
   wrong=copy.deepcopy(data);wrong['template_sha256']='wrong'
   try:t.build(svg,wrong,output);raise AssertionError('Wrong template accepted')
   except ValueError:pass
  t.build_pages(svg,{'_base':str(td),'pages':[data,data]},td/'multi.pptx');p=Presentation(td/'multi.pptx');assert len(p.slides)==2 and all(any(z.shape_type==13 for z in sl.shapes) for sl in p.slides);passed.append('multi-page native shapes and image relationships preserved')
 s.dump(root/'assets/strict-test-report.json',{'passed':len(passed),'checks':passed,'scope':'structure and behavior; visual review separate'});print('\n'.join('PASS '+x for x in passed))
if __name__=='__main__':main()
