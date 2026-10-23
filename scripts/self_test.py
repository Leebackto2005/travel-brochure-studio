"""Behavioral tests: output, exact edits, rejected mistakes and crop geometry."""
import copy, tempfile, zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
import studio as s

def main():
 passed=[]
 with tempfile.TemporaryDirectory(prefix='brochure-test-') as td:
  root=Path(td);d=s.specimen('editorial');d['_base']=td
  assert not s.validate(d)['errors'];passed.append('valid specimen')
  for label,mutate in [
   ('unknown image',lambda x:x['pages'][0].update(asset='bogus')),
   ('missing image',lambda x:x['assets'][0].update(path='missing.jpg')),
   ('capacity',lambda x:x['pages'][4]['days'].append(copy.deepcopy(x['pages'][4]['days'][0]))),
   ('numeric mismatch',lambda x:x['pages'][5]['rows'][0].update(zh='20KG',en='23KG')),
   ('room pairing',lambda x:x['pages'][3]['items'][0].pop('room'))]:
   bad=copy.deepcopy(d);mutate(bad);assert s.validate(bad)['errors'],label;passed.append('reject '+label)
  Image.new('RGB',(1200,400),'red').save(root/'wide.png');Image.new('RGB',(400,1200),'blue').save(root/'tall.png')
  d['assets'][0].update(path='wide.png',focus=[1,0]);d['assets'][1].update(path='tall.png',fit='contain')
  for style in s.STYLES:
   d['style']=style;out=root/(style+'.pptx');s.Painter(d).build(out);p=Presentation(out);assert len(p.slides)==6
   texts=[z.text for sl in p.slides for z in sl.shapes if z.has_text_frame];assert d['pages'][4]['days'][0]['zh'] in texts
   pics=[z for z in p.slides[0].shapes if z.shape_type==13];assert len(pics)==2
   c=pics[0];src=3;effective=src*(1-c.crop_left-c.crop_right)/(1-c.crop_top-c.crop_bottom);assert abs(effective-c.width/c.height)<.001
   assert abs(pics[1].width/pics[1].height-1/3)<.001
   passed.append('editable '+style+' and aspect-correct crops')
  src=root/'editorial.pptx';p=Presentation(src);sh=next(z for z in p.slides[0].shapes if z.has_text_frame and z.text=='海岛慢旅行')
  # Simulate mixed-format original text across two runs.
  par=sh.text_frame.paragraphs[0];par.clear();par.add_run().text='海岛';r=par.add_run();r.text='慢旅行';r.font.bold=True;p.save(src)
  edits=root/'edits.json';s.dump(edits,{'changes':[{'slide':1,'shape_id':sh.shape_id,'old':'岛慢旅','new':'岛好旅'}]});dst=root/'patched.pptx';s.patch(src,edits,dst)
  with zipfile.ZipFile(src) as a,zipfile.ZipFile(dst) as b:
   changed=[n for n in a.namelist() if a.read(n)!=b.read(n)];assert changed==['ppt/slides/slide1.xml']
  assert any(z.has_text_frame and z.text=='海岛好旅行' for z in Presentation(dst).slides[0].shapes);passed.append('cross-run exact edit; every unrelated ZIP member unchanged')
  try:s.patch(src,edits,src);raise AssertionError('overwrite allowed')
  except ValueError:passed.append('source overwrite rejected')
  s.dump(edits,{'changes':[{'slide':1,'shape_id':sh.shape_id,'old':'不存在','new':'替换'}]})
  try:s.patch(src,edits,dst);raise AssertionError('missing target allowed')
  except ValueError:passed.append('unmatched edit rejected')
  assert s.verify(src,root/'absent')['errors'];passed.append('unrendered deck cannot pass verification')
  before=root/'before';after=root/'after';before.mkdir();after.mkdir();Image.new('RGB',(10,10),'red').save(before/'P01.png');Image.new('RGB',(10,10),'blue').save(after/'P01.png')
  assert s.compare(before,after,set())['errors'];assert not s.compare(before,after,{1})['errors'];passed.append('authorized and unauthorized pixel changes differentiated')
 print('\n'.join('PASS '+x for x in passed));s.dump(s.ROOT/'assets'/'self-test-report.json',{'passed':len(passed),'checks':passed,'scope':'script behavior only; visual inspection separate'})
if __name__=='__main__':main()
