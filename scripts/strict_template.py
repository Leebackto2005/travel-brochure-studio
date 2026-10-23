"""Fill supplied SVG geometry without reflow; convert these four templates to native PPTX."""
import argparse, copy, io, json, math, re, sys, tempfile
from pathlib import Path
from lxml import etree as ET
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
import studio as s

EMU=12700
def tag(e):return ET.QName(e).localname
def val(e,k,default=0):return float(e.get(k,default))
def measure(text,size,serif=False,bold=False):
 if serif and text:
  western=Path('C:/Windows/Fonts')/('georgiab.ttf' if bold else 'georgia.ttf')
  chinese=Path('C:/Windows/Fonts/simsun.ttc')
  if western.exists() and chinese.exists():
   wf=ImageFont.truetype(str(western),max(1,round(size*4)));cf=ImageFont.truetype(str(chinese),max(1,round(size*4)))
   return sum((wf if ord(ch)<256 else cf).getlength(ch) for ch in text)/4
 path=Path('C:/Windows/Fonts')/('simsun.ttc' if serif else 'msyhbd.ttc' if bold else 'msyh.ttc')
 if not path.exists():raise ValueError('Install Chinese fonts or configure font mapping before strict-template rendering')
 font=ImageFont.truetype(str(path),max(1,round(size*4)))
 return font.getlength(text)/4
def catalog(svg):
 root=ET.parse(str(svg)).getroot();w,h=map(float,root.get('viewBox').split()[2:]);els=list(root);texts=[];photos=[];bars=[]
 for i,e in enumerate(els):
  if tag(e)=='text':
   fs=val(e,'font-size');t=''.join(e.itertext());serif='Noto Serif' in e.get('font-family','');bold=val(e,'font-weight',400)>=600
   width=(measure(t,fs,serif,bold)+max(0,len(t)-1)*val(e,'letter-spacing'))*1.28+8
   x=val(e,'x');anchor=e.get('text-anchor','start');left=x-width if anchor=='end' else x-width/2 if anchor=='middle' else x
   texts.append(dict(id=f'text-{len(texts)+1:02d}',element=i,text=t,x=left,y=val(e,'y')-fs*.91,width=width,height=fs*1.7,font_size=fs,anchor=anchor,baseline=val(e,'y')))
  if tag(e)=='rect' and e.get('stroke-width')=='1.15' and val(e,'width')>100 and val(e,'height')>70:
   photos.append(dict(id=f'photo-{len(photos)+1:02d}',element=i,x=val(e,'x'),y=val(e,'y'),width=val(e,'width'),height=val(e,'height')))
  if tag(e)=='rect' and 6<=val(e,'height')<=10 and val(e,'rx')>=3:bars.append((i,e))
 groups=[]
 for i,e in bars:
  if groups and val(e,'x')==val(groups[-1][-1][1],'x') and 0<val(e,'y')-val(groups[-1][-1][1],'y')<=26:groups[-1].append((i,e))
  else:groups.append([(i,e)])
 bodies=[]
 for g in groups:
  x=val(g[0][1],'x');y=val(g[0][1],'y');bottom=val(g[-1][1],'y')+val(g[-1][1],'height')
  bw=max(val(e,'width') for i,e in g);bh=max(18,bottom-y+11);bg='FFFFFF'
  for e in els[:g[0][0]]:
   if tag(e)=='rect' and e.get('fill','').startswith('#') and val(e,'x')<=x and val(e,'y')<=y and val(e,'x')+val(e,'width')>=x+bw and val(e,'y')+val(e,'height')>=y+bh-3:bg=e.get('fill')[1:]
  if len(bg)==3:bg=''.join(c*2 for c in bg)
  light=sum(int(bg[i:i+2],16)*weight for i,weight in zip((0,2,4),(.2126,.7152,.0722)))
  bodies.append(dict(id=f'body-{len(bodies)+1:02d}',elements=[i for i,e in g],x=x,y=y-3,width=bw,height=bh,font_size=14,color='FFFFFF' if light<110 else '192934'))
 return dict(template=str(Path(svg).resolve()),sha256=s.sha(svg),width=w,height=h,texts=texts,photos=photos,bodies=bodies,font_mapping={'Noto Sans CJK SC':'Microsoft YaHei','Noto Serif CJK SC':'SimSun (CJK) / Georgia (Latin)','note':'Original fonts absent on this machine; glyph appearance differs; geometry remains fixed'})
def paint(shape,e,scale):
 shape.shadow.inherit=False
 color=e.get('fill','none');opacity=val(e,'opacity',1)
 if color=='none':shape.fill.background()
 elif color.startswith('#'):
  if len(color)==4:color='#'+''.join(c*2 for c in color[1:])
  shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color[1:])
  if opacity<1:
   alpha=OxmlElement('a:alpha');alpha.set('val',str(round(opacity*100000)));shape.fill._xPr.solidFill.srgbClr.append(alpha)
 else:raise ValueError('Unsupported fill: '+color)
 stroke=e.get('stroke','none')
 if stroke=='none':shape.line.fill.background()
 else:
  if len(stroke)==4:stroke='#'+''.join(c*2 for c in stroke[1:])
  shape.line.color.rgb=RGBColor.from_string(stroke[1:]);shape.line.width=Pt(val(e,'stroke-width',1)*scale)
  if opacity<1:
   alpha=OxmlElement('a:alpha');alpha.set('val',str(round(opacity*100000)));shape.line._get_or_add_ln().solidFill.srgbClr.append(alpha)
def path_points(d):
 tokens=re.findall(r'[MLlZz]|[-+]?(?:\d*\.\d+|\d+)',d);pts=[];x=y=0;i=0;closed=False
 while i<len(tokens):
  cmd=tokens[i];i+=1
  if cmd.upper()=='Z':closed=True;continue
  if cmd not in ('M','L','l'):raise ValueError('Unsupported SVG path command')
  a,b=float(tokens[i]),float(tokens[i+1]);i+=2
  x,y=(x+a,y+b) if cmd=='l' else (a,b);pts.append((x,y))
 return pts,closed
def text_shape(slide,box,text,scale,color='192934',serif=False,bold=False,spacing=0):
 shape=slide.shapes.add_textbox(int(box['x']*scale*EMU),int(box['y']*scale*EMU),int(box['width']*scale*EMU),int(box['height']*scale*EMU));shape.name=box['id'];tf=shape.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for n,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if n==0 else tf.add_paragraph();p.text=line;p.font.name='Georgia' if serif else 'Microsoft YaHei';p.font.size=Pt(box['font_size']*scale);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.line_spacing=1.05
  p.alignment={'end':PP_ALIGN.RIGHT,'middle':PP_ALIGN.CENTER}.get(box.get('anchor'),PP_ALIGN.LEFT)
  for r in p.runs:
   rp=r._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface','SimSun' if serif else 'Microsoft YaHei');rp.append(ea)
   if spacing:rp.set('spc',str(round(spacing*scale*100)))
 return shape
def check_fit(t,b,serif=False,bold=False,spacing=0):
 fs=b['font_size'];lines=0
 for para in t.split('\n'):
  width=0;lines+=1
  for ch in para:
   cw=measure(ch,fs,serif,bold)+spacing
   if width+cw>b['width']-2:lines+=1;width=0
   width+=cw
 if lines*fs*1.05>b['height']+1:raise ValueError(f'{b["id"]}: content exceeds template slot; use continuation page or approved summary')
def build(svg,data,out):
 cat=catalog(svg);data=data or {};expected=data.get('template_sha256')
 if expected and expected!=cat['sha256']:raise ValueError('Template changed; rebuild slot mapping')
 if Path(svg).resolve()==Path(out).resolve():raise ValueError('Do not overwrite source')
 root=ET.parse(str(svg)).getroot();els=list(root);scale=841.89/cat['height'];prs=Presentation();prs.slide_width=int(cat['width']*scale*EMU);prs.slide_height=int(cat['height']*scale*EMU);slide=prs.slides.add_slide(prs.slide_layouts[6])
 texts={b['element']:b for b in cat['texts']};bodymap={b['id']:b for b in cat['bodies']};photomap={b['id']:b for b in cat['photos']};repl=data.get('texts',{});bodydata=data.get('bodies',{});photodata=data.get('photos',{});skip=set();bodyfirst={};log=[]
 unknown=set(repl)-{b['id'] for b in cat['texts']} | (set(bodydata)-set(bodymap)) | (set(photodata)-set(photomap))
 if unknown:raise ValueError('Unknown slots: '+str(unknown))
 for ident,t in bodydata.items():
  b=bodymap[ident];check_fit(t,b);skip.update(b['elements']);bodyfirst[b['elements'][0]]=(b,t)
 # Remove placeholder labels only when replacing that exact photo frame.
 for ident in photodata:
  b=photomap[ident]
  for i,t in texts.items():
   if b['x']<=t['x']<b['x']+b['width'] and b['y']<=t['baseline']<=b['y']+b['height']:skip.add(i)
 for i,e in enumerate(els):
  if i in bodyfirst:
   b,t=bodyfirst[i];text_shape(slide,b,t,scale,b.get('color','192934'))
  if i in skip:continue
  kind=tag(e)
  if kind=='defs':continue
  if kind=='text':
   b=texts[i];original=b['text'];t=repl.get(b['id'],original)
   if isinstance(t,dict):
    if t.get('old')!=original:raise ValueError('Original text changed for '+b['id'])
    t=t['new']
   serif='Noto Serif' in e.get('font-family','');bold=val(e,'font-weight',400)>=600;spacing=val(e,'letter-spacing')
   if t!=original:check_fit(t,b,serif,bold,spacing)
   color=e.get('fill','#192934')[1:]
   if len(color)==3:color=''.join(c*2 for c in color)
   text_shape(slide,b,t,scale,color,serif,bold,spacing);continue
  if kind=='rect':
   x,y,w,h=[val(e,k) for k in ('x','y','width','height')];shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if val(e,'rx') else MSO_SHAPE.RECTANGLE,int(x*scale*EMU),int(y*scale*EMU),int(w*scale*EMU),int(h*scale*EMU))
   if val(e,'rx'):shape.adjustments[0]=min(.5,val(e,'rx')/max(1,min(w,h)))
  elif kind=='circle':
   r=val(e,'r');shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,int((val(e,'cx')-r)*scale*EMU),int((val(e,'cy')-r)*scale*EMU),int(2*r*scale*EMU),int(2*r*scale*EMU))
  elif kind=='path':
   pts,closed=path_points(e.get('d'));fb=slide.shapes.build_freeform(pts[0][0],pts[0][1],scale=scale*EMU);fb.add_line_segments(pts[1:],close=closed);shape=fb.convert_to_shape()
  else:raise ValueError('Unsupported SVG element: '+kind)
  shape.name=f'svg-element-{i}';paint(shape,e,scale)
 # Native pictures exactly match the supplied frame bounds.
 painter=s.Painter({'style':'editorial','assets':[],'pages':[{}],'_base':data.get('_base','.')});painter.s=slide
 for ident,a in photodata.items():
  b=photomap[ident];a={'id':ident,'source':'用户配图','verified':False,**a};painter.assets[ident]=a;painter.photo(ident,b['x']*scale/72,b['y']*scale/72,b['width']*scale/72,b['height']*scale/72)
  log.append(dict(slot=ident,box={k:b[k] for k in ('x','y','width','height')},source=a['source'],fit=a.get('fit','cover'),focus=a.get('focus',[.5,.5])))
 Path(out).parent.mkdir(parents=True,exist_ok=True);prs.save(out);s.dump(str(out)+'.receipt.json',{'template_sha256':cat['sha256'],'template_size':[cat['width'],cat['height']],'geometry':'Original SVG positions and sizes retained; no reflow','font_mapping':cat['font_mapping'],'text_changes':repl,'body_changes':bodydata,'photos':log,'visual_review':'required'});return cat
def combine(inputs,out):
 """Copy native shapes and their embedded images; reject mixed canvas sizes."""
 result=Presentation();dims=None
 for path in inputs:
  deck=Presentation(path);size=(deck.slide_width,deck.slide_height)
  if dims and size!=dims:raise ValueError('Cannot combine different template dimensions')
  dims=size;result.slide_width,result.slide_height=size
  for old in deck.slides:
   new=result.slides.add_slide(result.slide_layouts[6])
   for shape in old.shapes:
    el=copy.deepcopy(shape.element)
    for blip in el.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
     key='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed';rid=blip.get(key)
     if rid:
      part=old.part.related_part(rid);_,newrid=new.part.get_or_add_image_part(io.BytesIO(part.blob));blip.set(key,newrid)
    new.shapes._spTree.insert_element_before(el,'p:extLst')
 result.save(out)
def build_pages(svg,data,out):
 pages=data.get('pages')
 if pages is None:return build(svg,data,out)
 if not pages:raise ValueError('pages must not be empty')
 with tempfile.TemporaryDirectory(prefix='strict-template-') as td:
  paths=[];receipts=[]
  for i,page in enumerate(pages):
   p=Path(td)/f'{i}.pptx';build(svg,{'_base':data.get('_base','.'),'template_sha256':data.get('template_sha256'),**page},p);paths.append(p);receipts.append(s.read(str(p)+'.receipt.json'))
  Path(out).parent.mkdir(parents=True,exist_ok=True);combine(paths,out);s.dump(str(out)+'.receipt.json',{'pages':receipts,'pagination':'explicit, same original template geometry'})
def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('command',choices=['catalog','build']);p.add_argument('template');p.add_argument('--data');p.add_argument('--out',required=True);a=p.parse_args()
 if a.command=='catalog':s.dump(a.out,catalog(a.template))
 else:
  data=s.read(a.data) if a.data else {};data['_base']=str(Path(a.data).resolve().parent) if a.data else '.';build_pages(a.template,data,a.out)
 print(a.command+': complete')
if __name__=='__main__':main()
